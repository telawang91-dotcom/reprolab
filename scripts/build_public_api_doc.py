from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(
    r"D:\Desktop\深大A赛道\reprolab\最终提交材料（本地勿提交）"
    r"\02-Docker部署与远程调用\03-ReproLab公网可调用接口说明.docx"
)
BASE_URL = "https://reverse-notebooks-trigger-rogers.trycloudflare.com"
PROJECT_ID = "4713147e-4ac2-4c4d-a01f-e748139dc506"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "172B4D"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
GREEN = "16794A"
WHITE = "FFFFFF"
BORDER = "CBD5E1"


def set_run_font(
    run,
    *,
    name: str = "Calibri",
    east_asia: str = "Microsoft YaHei",
    size: float | None = None,
    bold: bool | None = None,
    color: str | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_paragraph_border(paragraph, color: str = BORDER) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "4")
        tag.set(qn("w:color"), color)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run_props.extend([fonts, color, underline])
    run.append(run_props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = paragraph.add_run("第 ")
    set_run_font(label, size=9, color=MUTED)
    for field_name in ("PAGE",):
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = field_name
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        text = OxmlElement("w:t")
        text.text = "1"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run = OxmlElement("w:r")
        run.extend([begin, instruction, separate, text, end])
        paragraph._p.append(run)
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, size=9, color=MUTED)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("ReproLab  |  公网 REST API")
    set_run_font(run, size=9, color=MUTED, bold=True)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    add_page_number(footer)


def add_title_block(doc: Document) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(3)
    run = kicker.add_run("外部评审调用指南")
    set_run_font(run, size=10, color=GREEN, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(5)
    run = title.add_run("ReproLab 可调用接口说明")
    set_run_font(run, size=25, color=INK, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("公网 REST API · 同步与异步任务调用 · 溯源产物返回")
    set_run_font(run, size=12.5, color=MUTED)

    status = doc.add_table(rows=1, cols=1)
    set_table_geometry(status, [9360])
    cell = status.cell(0, 0)
    set_cell_shading(cell, "EAF7F0")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    lead = p.add_run("当前状态：")
    set_run_font(lead, size=10.5, bold=True, color=GREEN)
    detail = p.add_run("公网接口在线；当前无需 Token；建议长任务优先使用异步接口。")
    set_run_font(detail, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_metadata(doc: Document) -> None:
    table = doc.add_table(rows=5, cols=2)
    set_table_geometry(table, [1700, 7660])
    rows = [
        ("Base URL", BASE_URL),
        ("Swagger", f"{BASE_URL}/docs"),
        ("健康检查", f"{BASE_URL}/health"),
        ("项目 ID", PROJECT_ID),
        ("鉴权", "当前无需 Authorization 请求头"),
    ]
    for index, (label, value) in enumerate(rows):
        left, right = table.rows[index].cells
        set_cell_shading(left, LIGHT_BLUE)
        left_p = left.paragraphs[0]
        left_p.paragraph_format.space_after = Pt(0)
        set_run_font(left_p.add_run(label), size=10, bold=True, color=DARK_BLUE)
        right_p = right.paragraphs[0]
        right_p.paragraph_format.space_after = Pt(0)
        if value.startswith("https://"):
            add_hyperlink(right_p, value, value)
        else:
            set_run_font(right_p.add_run(value), size=10, color=INK)


def add_endpoint_table(doc: Document) -> None:
    doc.add_heading("1. 接口一览", level=1)
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [1450, 3400, 4510])
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell, text in zip(header.cells, ("方法", "路径", "用途"), strict=True):
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(text), size=9.5, bold=True, color=DARK_BLUE)
    data = [
        ("GET", "/health", "检查服务和数据库是否在线"),
        ("POST", "/api/v1/agent/jobs", "异步提交任务（推荐）"),
        ("GET", "/api/v1/agent/jobs/{job_id}", "查询异步任务状态与结果"),
        ("DELETE", "/api/v1/agent/jobs/{job_id}", "取消尚未完成的异步任务"),
        ("POST", "/api/v1/agent/invoke", "同步执行短任务并直接返回结果"),
    ]
    for method, path, purpose in data:
        cells = table.add_row().cells
        for index, text in enumerate((method, path, purpose)):
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if index == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_run_font(p.add_run(text), name="Consolas", east_asia="Microsoft YaHei", size=9, bold=True, color=GREEN)
            elif index == 1:
                set_run_font(p.add_run(text), name="Consolas", east_asia="Microsoft YaHei", size=8.6, color=INK)
            else:
                set_run_font(p.add_run(text), size=9.5, color=INK)
    set_table_geometry(table, [1450, 3400, 4510])


def add_request_fields(doc: Document) -> None:
    doc.add_heading("2. 通用请求参数", level=1)
    lead = doc.add_paragraph("同步和异步接口使用相同的 JSON 请求体。")
    lead.paragraph_format.space_after = Pt(7)

    table = doc.add_table(rows=1, cols=4)
    set_table_geometry(table, [2200, 1450, 1150, 4560])
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell, text in zip(header.cells, ("字段", "类型", "必填", "说明"), strict=True):
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(text), size=9.5, bold=True, color=DARK_BLUE)
    fields = [
        ("project_id", "UUID", "是", f"使用本文提供的项目 ID：{PROJECT_ID}"),
        ("task", "string", "是", "自然语言任务，长度 1–10,000 字符"),
        ("inputs.dataset_ids", "UUID[]", "否", "需要分析已上传数据时传入数据集 ID；无数据时传空数组"),
        ("inputs.skill_id", "UUID/null", "否", "指定可选技能；不指定时可省略"),
    ]
    for field, kind, required, detail in fields:
        cells = table.add_row().cells
        for index, text in enumerate((field, kind, required, detail)):
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if index in (1, 2):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(
                p.add_run(text),
                name="Consolas" if index == 0 else "Calibri",
                east_asia="Microsoft YaHei",
                size=8.8 if index == 0 else 9.2,
                bold=index == 2,
                color=INK,
            )
    set_table_geometry(table, [2200, 1450, 1150, 4560])


def add_code_block(doc: Document, code: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.08)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = True
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shading)
    set_paragraph_border(paragraph, BORDER)
    lines = code.splitlines()
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        set_run_font(
            run,
            name="Consolas",
            east_asia="Microsoft YaHei",
            size=8.25,
            color="1F2937",
        )
        if index < len(lines) - 1:
            run.add_break()


def add_step(doc: Document, number: str, title: str, detail: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    badge = paragraph.add_run(f"{number}  ")
    set_run_font(badge, size=11, color=GREEN, bold=True)
    label = paragraph.add_run(title)
    set_run_font(label, size=11, color=INK, bold=True)
    if detail:
        desc = paragraph.add_run(f"  {detail}")
        set_run_font(desc, size=10.5, color=MUTED)


def add_async_section(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("3. 推荐调用方式：异步任务", level=1)
    intro = doc.add_paragraph(
        "异步方式会先返回 job_id，调用方随后轮询任务状态。该方式适合需要生成代码、启动 Python 内核或执行数据分析的任务。"
    )
    intro.paragraph_format.space_after = Pt(4)

    add_step(doc, "步骤 1", "提交任务", "使用唯一 Idempotency-Key 防止网络重试造成重复任务。")
    add_code_block(
        doc,
        f"""curl -X POST "{BASE_URL}/api/v1/agent/jobs" \\
  -H "Content-Type: application/json" \\
  -H "Idempotency-Key: evaluator-test-001" \\
  -d '{{
    "project_id": "{PROJECT_ID}",
    "task": "请使用 Python 计算 8 + 9，并清楚给出结果。",
    "inputs": {{
      "dataset_ids": []
    }}
  }}'""",
    )

    add_step(doc, "步骤 2", "读取受理响应", "保存 job_id，并使用 status_url 查询状态。")
    add_code_block(
        doc,
        """{
  "job_id": "7fad80ae-35eb-4229-ab8e-df8a5b128c01",
  "status": "queued",
  "status_url": "/api/v1/agent/jobs/7fad80ae-35eb-4229-ab8e-df8a5b128c01",
  "created_at": "2026-07-30T09:25:00Z"
}""",
    )

    add_step(doc, "步骤 3", "轮询任务状态", "建议每 2–3 秒查询一次，直到进入终态。")
    add_code_block(
        doc,
        f"""curl "{BASE_URL}/api/v1/agent/jobs/{{job_id}}\"""",
    )

    note = doc.add_table(rows=1, cols=1)
    set_table_geometry(note, [9360])
    cell = note.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run("终态："), size=10, bold=True, color=DARK_BLUE)
    set_run_font(
        p.add_run("succeeded（成功）、failed（失败）或 cancelled（已取消）。"),
        size=10,
        color=INK,
    )


def add_result_and_sync(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("4. 获取任务结果", level=1)
    p = doc.add_paragraph(
        "当 status 为 succeeded 时，result 字段包含自然语言回答、可信产物、血缘信息和核验报告。"
    )
    p.paragraph_format.space_after = Pt(4)
    add_code_block(
        doc,
        """{
  "job_id": "7fad80ae-35eb-4229-ab8e-df8a5b128c01",
  "status": "succeeded",
  "result": {
    "result": "分析结果正文",
    "artifacts": [
      {
        "kind": "number",
        "value_json": {
          "value": 17,
          "mime_type": "application/vnd.reprolab.artifact+json"
        },
        "anchor": "⟦art_xxxx⟧"
      }
    ],
    "lineage": {},
    "verify_report": {
      "verdict": "pass",
      "items": []
    }
  }
}""",
    )

    doc.add_heading("5. 同步调用", level=1)
    p = doc.add_paragraph(
        "短任务可使用同步接口直接等待完整结果。复杂分析仍建议使用异步任务，避免中间网络超时。"
    )
    p.paragraph_format.space_after = Pt(4)
    add_code_block(
        doc,
        f"""curl -X POST "{BASE_URL}/api/v1/agent/invoke" \\
  -H "Content-Type: application/json" \\
  -d '{{
    "project_id": "{PROJECT_ID}",
    "task": "请使用 Python 计算 8 + 9，并清楚给出结果。",
    "inputs": {{
      "dataset_ids": []
    }}
  }}'""",
    )

    doc.add_heading("6. Python 调用示例", level=1)
    add_code_block(
        doc,
        f"""import requests

base_url = "{BASE_URL}"
payload = {{
    "project_id": "{PROJECT_ID}",
    "task": "请使用 Python 计算 8 + 9，并清楚给出结果。",
    "inputs": {{"dataset_ids": []}},
}}

response = requests.post(
    f"{{base_url}}/api/v1/agent/jobs",
    json=payload,
    headers={{"Idempotency-Key": "evaluator-test-001"}},
    timeout=30,
)
response.raise_for_status()
print(response.json())""",
    )


def add_notes(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("7. 返回内容说明", level=1)
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [2200, 7160])
    for cell, text in zip(table.rows[0].cells, ("字段", "含义"), strict=True):
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(text), size=9.5, bold=True, color=DARK_BLUE)
    rows = [
        ("result.result", "面向用户的最终自然语言回答"),
        ("result.artifacts", "执行过程中登记的数字、表格、图片等可信产物"),
        ("result.lineage", "产物与代码、运行和输入数据之间的溯源关系"),
        ("result.verify_report", "对数字、引用和产物一致性的核验结果"),
        ("error", "任务失败时返回的错误信息"),
    ]
    for field, detail in rows:
        cells = table.add_row().cells
        p1 = cells[0].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        set_run_font(p1.add_run(field), name="Consolas", east_asia="Microsoft YaHei", size=8.8, color=INK)
        p2 = cells[1].paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        set_run_font(p2.add_run(detail), size=9.5, color=INK)
    set_table_geometry(table, [2200, 7160])

    doc.add_heading("8. 在线验证结果", level=1)
    validation = doc.add_table(rows=4, cols=2)
    set_table_geometry(validation, [2500, 6860])
    checks = [
        ("健康检查", "HTTP 200，数据库 online"),
        ("Swagger / OpenAPI", "HTTP 200，可查看完整接口结构"),
        ("异步任务", "queued → succeeded"),
        ("真实计算", "8 + 9 返回 17，并生成 1 个溯源产物"),
    ]
    for index, (label, result) in enumerate(checks):
        left, right = validation.rows[index].cells
        if index % 2 == 0:
            set_cell_shading(left, CALLOUT)
            set_cell_shading(right, CALLOUT)
        p1 = left.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        set_run_font(p1.add_run(label), size=9.5, bold=True, color=DARK_BLUE)
        p2 = right.paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        set_run_font(p2.add_run(result), size=9.5, color=INK)

    doc.add_heading("9. 临时公网地址说明", level=1)
    note = doc.add_table(rows=1, cols=1)
    set_table_geometry(note, [9360])
    cell = note.cell(0, 0)
    set_cell_shading(cell, "FFF8E7")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run("重要："), size=10, bold=True, color="7A5A00")
    set_run_font(
        p.add_run(
            "该公网地址由本机临时隧道提供。调用期间必须保持电脑不休眠、Docker Desktop 正常运行且网络在线；"
            "重启、断网或停止隧道后地址会失效。当前接口无需 Token，请勿上传敏感资料。"
        ),
        size=10,
        color=INK,
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run("接口文档："), size=10, bold=True, color=DARK_BLUE)
    add_hyperlink(p, f"{BASE_URL}/docs", f"{BASE_URL}/docs")


def audit_document(doc: Document) -> None:
    section = doc.sections[0]
    assert section.page_width == Inches(8.5)
    assert section.page_height == Inches(11)
    assert section.left_margin == Inches(1)
    assert section.right_margin == Inches(1)
    for table in doc.tables:
        grid = table._tbl.tblGrid
        widths = [int(item.get(qn("w:w"))) for item in grid]
        assert sum(widths) == 9360, widths
        for row in table.rows:
            assert len(row.cells) == len(widths)
            for index, cell in enumerate(row.cells):
                tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
                assert tc_w is not None
                assert int(tc_w.get(qn("w:w"))) == widths[index]


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style_document(doc)
    add_title_block(doc)
    add_metadata(doc)
    add_endpoint_table(doc)
    add_request_fields(doc)
    add_async_section(doc)
    add_result_and_sync(doc)
    add_notes(doc)
    audit_document(doc)
    core = doc.core_properties
    core.title = "ReproLab 可调用接口说明"
    core.subject = "公网 REST API 外部调用指南"
    core.author = "ReproLab"
    core.keywords = "ReproLab, REST API, 异步任务, 溯源"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
